from flask import Flask,render_template,request,redirect,url_for,session,flash
import mysql.connector
from werkzeug.security import generate_password_hash,check_password_hash
import os,re
from werkzeug.utils import secure_filename
from pypdf import PdfReader
from docx import Document
from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(__file__),".env"))

app=Flask(__name__,template_folder="../frontend/templates",static_folder="../frontend/static")
app.secret_key="ai_interview_secret_key"

UPLOAD_FOLDER=os.path.join(os.path.dirname(__file__),"uploads")
ALLOWED_EXTENSIONS={"pdf","docx"}
os.makedirs(UPLOAD_FOLDER,exist_ok=True)
app.config["UPLOAD_FOLDER"]=UPLOAD_FOLDER

def get_db():
    return mysql.connector.connect(
        host="localhost",
        user="root",
        password=os.getenv("DB_PASSWORD"),
        database="ai_interview",
        port=3306
    )

def allowed_file(filename):
    return "." in filename and filename.rsplit(".",1)[1].lower() in ALLOWED_EXTENSIONS

def extract_pdf_text(file_path):
    text=""
    try:
        reader=PdfReader(file_path)
        for page in reader.pages:
            page_text=page.extract_text()
            if page_text:
                text+=page_text+" "
    except:
        return ""
    return text

def extract_docx_text(file_path):
    text=""
    try:
        document=Document(file_path)
        for paragraph in document.paragraphs:
            text+=paragraph.text+" "
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text+=cell.text+" "
    except:
        return ""
    return text

def extract_resume_text(file_path):
    extension=file_path.rsplit(".",1)[1].lower()
    if extension=="pdf":
        return extract_pdf_text(file_path)
    if extension=="docx":
        return extract_docx_text(file_path)
    return ""

def analyze_resume(text):
    text=text.lower()
    text=re.sub(r"\s+"," ",text).strip()
    resume_sections={
        "education":["education","academic","qualification","degree","bachelor","master","university","college"],
        "skills":["skills","technical skills","programming","python","java","javascript","html","css","sql","mysql","flask","machine learning","data science"],
        "experience":["experience","work experience","employment","internship","intern","developer","engineer","analyst"],
        "projects":["projects","project","developed","implemented","application","system"],
        "profile":["objective","summary","profile","career objective","professional summary"],
        "contact":["email","phone","mobile","linkedin","github","address"]
    }
    scores={}
    for section,keywords in resume_sections.items():
        matches=sum(1 for keyword in keywords if keyword in text)
        scores[section]=matches
    bill_words=[
        "invoice","bill","gst","gstin","tax invoice","invoice number",
        "customer invoice","amount payable","subtotal","total amount",
        "quantity","unit price","product details","warranty","serial number",
        "purchase invoice","cash memo","receipt","order number"
    ]
    bill_matches=sum(1 for word in bill_words if word in text)
    resume_words=[
        "resume","curriculum vitae","career objective","professional summary",
        "education","skills","experience","projects","certification",
        "internship","work experience","technical skills","achievement",
        "qualification","languages","profile"
    ]
    resume_matches=sum(1 for word in resume_words if word in text)
    if len(text)<100:
        return {
            "valid":False,
            "score":0,
            "skills":"",
            "education":"",
            "experience":"",
            "projects":""
        }
    if bill_matches>=2 and resume_matches<3:
        return {
            "valid":False,
            "score":0,
            "skills":"",
            "education":"",
            "experience":"",
            "projects":""
        }
    strong_resume_score=0
    if scores["education"]>=1:
        strong_resume_score+=20
    if scores["skills"]>=2:
        strong_resume_score+=20
    if scores["experience"]>=1:
        strong_resume_score+=20
    if scores["projects"]>=1:
        strong_resume_score+=15
    if scores["profile"]>=1:
        strong_resume_score+=10
    if scores["contact"]>=2:
        strong_resume_score+=10
    if len(text)>=500:
        strong_resume_score+=5
    if resume_matches<3:
        return {
            "valid":False,
            "score":strong_resume_score,
            "skills":"",
            "education":"",
            "experience":"",
            "projects":""
        }
    return {
        "valid":True,
        "score":min(strong_resume_score,100),
        "skills":"Detected" if scores["skills"]>=2 else "",
        "education":"Detected" if scores["education"]>=1 else "",
        "experience":"Detected" if scores["experience"]>=1 else "",
        "projects":"Detected" if scores["projects"]>=1 else ""
    }

def get_resume(user_id):
    db=get_db()
    cursor=db.cursor(dictionary=True)
    cursor.execute("SELECT * FROM resumes WHERE user_id=%s ORDER BY id DESC LIMIT 1",(user_id,))
    resume=cursor.fetchone()
    cursor.close()
    db.close()
    return resume

def resume_eligible(user_id):
    resume=get_resume(user_id)
    if not resume:
        return False
    return float(resume["resume_score"] or 0)>=75

@app.route("/")
def index():
    return render_template("index.html")

@app.route("/login",methods=["GET","POST"])
def login():
    if request.method=="POST":
        email=request.form.get("email","").strip()
        password=request.form.get("password","")
        if not email or not password:
            flash("Please enter email and password.")
            return redirect(url_for("login"))
        db=get_db()
        cursor=db.cursor(dictionary=True)
        cursor.execute("SELECT * FROM users WHERE email=%s",(email,))
        user=cursor.fetchone()
        cursor.close()
        db.close()
        if user and check_password_hash(user["password"],password):
            session.clear()
            session["user_id"]=user["id"]
            session["user_name"]=user["first_name"]
            session["user_email"]=user["email"]
            session["resume_uploaded"]=False
            return redirect(url_for("resume_analysis"))
        flash("Invalid email or password.")
        return redirect(url_for("login"))
    return render_template("login.html")

@app.route("/register",methods=["GET","POST"])
def register():
    if request.method=="POST":
        first_name=request.form.get("first_name","").strip()
        last_name=request.form.get("last_name","").strip()
        email=request.form.get("email","").strip()
        password=request.form.get("password","")
        confirm_password=request.form.get("confirm_password","")
        if not first_name or not last_name or not email or not password:
            flash("Please fill all required fields.")
            return redirect(url_for("register"))
        if password!=confirm_password:
            flash("Passwords do not match.")
            return redirect(url_for("register"))
        db=get_db()
        cursor=db.cursor()
        cursor.execute("SELECT id FROM users WHERE email=%s",(email,))
        existing=cursor.fetchone()
        if existing:
            cursor.close()
            db.close()
            flash("Email already registered.")
            return redirect(url_for("register"))
        hashed_password=generate_password_hash(password)
        cursor.execute(
            "INSERT INTO users (first_name,last_name,email,password) VALUES (%s,%s,%s,%s)",
            (first_name,last_name,email,hashed_password)
        )
        db.commit()
        cursor.close()
        db.close()
        flash("Registration successful. Please login.")
        return redirect(url_for("login"))
    return render_template("register.html")

@app.route("/resume-analysis",methods=["GET","POST"])
def resume_analysis():
    if "user_id" not in session:
        return redirect(url_for("login"))
    if request.method=="POST":
        if "resume" not in request.files:
            flash("Please select a resume file.")
            return redirect(url_for("resume_analysis"))
        file=request.files["resume"]
        if file.filename=="":
            flash("Please select a resume file.")
            return redirect(url_for("resume_analysis"))
        if not allowed_file(file.filename):
            flash("Please upload a PDF or DOCX resume file.")
            return redirect(url_for("resume_analysis"))
        filename=secure_filename(file.filename)
        user_id=session["user_id"]
        user_folder=os.path.join(app.config["UPLOAD_FOLDER"],str(user_id))
        os.makedirs(user_folder,exist_ok=True)
        file_path=os.path.join(user_folder,filename)
        file.save(file_path)
        text=extract_resume_text(file_path)
        if not text:
            if os.path.exists(file_path):
                os.remove(file_path)
            session["resume_uploaded"]=False
            flash("Unable to read this file. Please upload a proper text-based resume.")
            return redirect(url_for("resume_analysis"))
        analysis=analyze_resume(text)
        db=get_db()
        cursor=db.cursor()
        cursor.execute("DELETE FROM resumes WHERE user_id=%s",(user_id,))
        cursor.execute(
            "INSERT INTO resumes (user_id,file_name,file_path,resume_score,skills,education,experience,projects) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",
            (
                user_id,
                filename,
                file_path,
                analysis["score"],
                analysis["skills"],
                analysis["education"],
                analysis["experience"],
                analysis["projects"]
            )
        )
        db.commit()
        cursor.close()
        db.close()
        if not analysis["valid"]:
            session["resume_uploaded"]=False
            flash("The uploaded file does not appear to be a valid resume. Please upload your resume.")
        elif analysis["score"]<75:
            session["resume_uploaded"]=True
            flash("Resume analyzed successfully, but the score is below 75%. Aptitude Round is locked.")
        else:
            session["resume_uploaded"]=True
            flash("Resume analyzed successfully. Aptitude Round is unlocked.")
        return redirect(url_for("resume_analysis"))
    resume=get_resume(session["user_id"])
    return render_template("resume_analysis.html",resume=resume)

@app.route("/aptitude-round")
def aptitude_round():
    if "user_id" not in session:
        return redirect(url_for("login"))
    if not session.get("resume_uploaded",False):
        flash("Please upload your resume before starting the Aptitude Round.")
        return redirect(url_for("resume_analysis"))
    resume=get_resume(session["user_id"])
    if not resume:
        session["resume_uploaded"]=False
        flash("Please upload your resume before starting the Aptitude Round.")
        return redirect(url_for("resume_analysis"))
    if float(resume["resume_score"] or 0)<75:
        flash("Your resume score must be 75% or above to start the Aptitude Round.")
        return redirect(url_for("resume_analysis"))
    return render_template("aptitude_round.html")

@app.route("/technical-round")
def technical_round():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("technical_round.html")

@app.route("/coding-skill")
def coding_skill():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("coding_skill.html")

@app.route("/personal-interview")
def personal_interview():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("personal_interview.html")

@app.route("/hr-interview")
def hr_interview():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("hr_interview.html")

@app.route("/mock-interview")
def mock_interview():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("mock_interview.html")

@app.route("/performance-tracking")
def performance_tracking():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("performance_tracking.html")

@app.route("/personalized-recommendations")
def personalized_recommendations():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("personalized_recommendations.html")

@app.route("/feedback")
def feedback():
    if "user_id" not in session:
        return redirect(url_for("login"))
    return render_template("feedback.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("index"))

if __name__=="__main__":
    app.run(debug=True)